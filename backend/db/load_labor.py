"""Install_Labor_HackOHIO.xlsx -> labor_* reference tables.

The workbook stacks several small titled tables per sheet, so blocks are found
by their "... Average Hours" title rows rather than read as one grid.
"""
import openpyxl
import pandas as pd
from sqlalchemy import text
from sqlalchemy.engine import Engine

from .config import DATA_DIR
from .sources import LABOR_FILE

# Header text in a block -> output column.
BLOCK_HEADERS = {
    "size group": "size_label",
    "size (kva)": "size_label",
    "height range": "size_range",
    "cu count": "cu_count",
    "avg install hours": "avg_install_hours",
    "avg removal hours": "avg_removal_hours",
}
GROUP_AVG_COLS = ["cu_group", "size_label", "size_range", "cu_count", "avg_install_hours", "avg_removal_hours"]


def _cell(ws, row, col):
    v = ws.cell(row=row, column=col).value
    return v.strip() if isinstance(v, str) else v


def parse_avg_blocks(ws, first_col: int, last_col: int) -> list[dict]:
    rows = []
    r = 1
    while r <= ws.max_row:
        title = _cell(ws, r, first_col)
        if not (isinstance(title, str) and "Average Hours" in title):
            r += 1
            continue
        group = title.split(" CU Average Hours")[0].split(" Average Hours")[0]
        header_cols = {}
        for c in range(first_col, last_col + 1):
            h = _cell(ws, r + 1, c)
            if isinstance(h, str) and h.lower() in BLOCK_HEADERS:
                header_cols[c] = BLOCK_HEADERS[h.lower()]
        r += 2
        # Data rows continue until a row with no numeric values (blank, note, or next title).
        while r <= ws.max_row:
            values = {col: _cell(ws, r, c) for c, col in header_cols.items()}
            if not any(isinstance(v, (int, float)) for v in values.values()):
                break
            rows.append({"cu_group": group, **values})
            r += 1
    return rows


def labor_cu_hours(wb) -> pd.DataFrame:
    ws = wb["CNC"]
    rows = []
    for r in range(2, ws.max_row + 1):
        cu = _cell(ws, r, 1)
        if cu:
            rows.append({"cd_cu": cu, "install_hours": _cell(ws, r, 2), "remove_hours": _cell(ws, r, 3)})
    return pd.DataFrame(rows)


def labor_cu_group_avg(wb) -> pd.DataFrame:
    rows = parse_avg_blocks(wb["Labor time material related"], 1, 4)
    rows += parse_avg_blocks(wb["CNC"], 5, 7)  # connector summary beside the CU list
    df = pd.DataFrame(rows).reindex(columns=GROUP_AVG_COLS)
    df["size_label"] = df["size_label"].map(lambda v: None if pd.isna(v) else str(v))
    return df


def labor_adders(wb) -> pd.DataFrame:
    ws = wb["Labor only work"]
    rows = []
    for r in range(2, ws.max_row + 1):
        intent = _cell(ws, r, 1)
        if not intent:
            continue
        b, c = _cell(ws, r, 2), _cell(ws, r, 3)
        # Rows 2-8 fill work_type (B) + description (C); later rows put the description in B.
        work_type, description = (b, c) if c else (None, b)
        rows.append({
            "original_intent": intent,
            "work_type": work_type,
            "description": description,
            "install_hours": _cell(ws, r, 4),
            "unit": _cell(ws, r, 5) or "per occurrence",
        })
    return pd.DataFrame(rows)


def labor_indirect_multiplier(wb) -> pd.DataFrame:
    ws = wb["Indirect Labor Multiplier"]
    return pd.DataFrame([{"name": _cell(ws, 1, 1), "value": _cell(ws, 1, 2), "note": _cell(ws, 3, 1)}])


def labor_cu_categories(wb) -> pd.DataFrame:
    ws = wb["CU Group descriptions"]
    rows = [
        {"category": _cell(ws, r, 1), "description": _cell(ws, r, 2)}
        for r in range(2, ws.max_row + 1)
        if _cell(ws, r, 1)
    ]
    return pd.DataFrame(rows)


def load(engine: Engine) -> None:
    wb = openpyxl.load_workbook(DATA_DIR / LABOR_FILE, data_only=True, read_only=False)
    tables = {
        "labor_cu_hours": labor_cu_hours(wb),
        "labor_cu_group_avg": labor_cu_group_avg(wb),
        "labor_adders": labor_adders(wb),
        "labor_indirect_multiplier": labor_indirect_multiplier(wb),
        "labor_cu_categories": labor_cu_categories(wb),
    }
    for name, df in tables.items():
        with engine.begin() as conn:
            conn.execute(text(f"DROP TABLE IF EXISTS {name} CASCADE"))
        df.to_sql(name, engine, index=False)
        print(f"  {name}: {len(df):,} rows")
